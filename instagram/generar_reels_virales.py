from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# (cuenta, seguidores, [(vistas, descripcion, tipo, sirve)])
# tipo: Humor, Personal, Educativo, Opinión, Pago, Repost, Venta
D = [
("saakshiityaagii",26500,[(2_000_000,'"How to Aura-Farm on LinkedIn" (sátira de la jerga de LinkedIn, 137K likes)',"Humor","Sí, molde"),(93_300,'"Aura farming on LinkedIn 101" (secuela)',"Humor","Sí, molde"),(72_400,"Resumen de noticias de Tech Twitter/X","Educativo","Parcial")]),
("diandraescobar",22000,[(354_000,"Aparece en un ranking junto a presidentes de empresas gracias a LinkedIn (en español)","Personal","Sí, molde"),(149_000,"Llamado a ayudar a Colombia (crisis)","Personal","No"),(47_900,'"Colombia needs our help": familias afectadas en Chocó',"Personal","No")]),
("charlotte.arsenault",230000,[(801_000,'"Layered girlies unite" (identidad multifacética)',"Humor","Sí, molde"),(781_000,'Serie "cómo entrenar a tu novio para sacarte fotos lindas", ep. 2',"Humor","Parcial"),(775_000,"Paso a paso de poses para fotos (audio karaoke)","Educativo","No")]),
("prettylittlemarketer",195000,[(1_600_000,"Semrush: herramienta de IA para ideas de contenido","Pago","No"),(1_400_000,'Snapchat: "add Snapchat to your media plan"',"Pago","No"),(692_000,"Snapdragon UK: unboxing de laptop + sorteo","Pago","No")]),
("womanceomindset",1000000,[(28_900_000,'"Besties for life — tag your besties" (repost de @amycowgillx, 2023)',"Repost","Parcial"),(13_100_000,'"No time for bad energy, tag your biz bestie" (2023)',"Repost","Parcial"),(9_000_000,'"Mood all 2023. No time for bad energy"',"Repost","No")]),
("thevaleriaramirez",55000,[(894_000,"Cómo arreglar el parpadeo de luces en videos (frame rate 50Hz)","Educativo","No"),(880_000,'"Todos se pasaron a Claude y vos seguís con ChatGPT"',"Opinión","Sí, molde"),(358_000,"+3.000 seguidores en una noche con Trial Reels (framework)","Educativo","Parcial")]),
("artofinfluenceco",408000,[(10_000_000,'"¿Querés ser creadora y trabajar con marcas? Training gratis en la bio" (2023)',"Venta","Parcial"),(8_800_000,"Tutorial: foto borrosa con Snapseed","Educativo","No"),(6_800_000,'"¿Qué pasaría si vas ALL IN en redes en 2024?" (venta de su programa)',"Venta","Sí, molde")]),
("britcrowther",10300,[(2_200_000,'"Never in my likes, but always in my views" (collab @memequeen)',"Humor","Sí, molde"),(1_600_000,'"Life\'s hard. Make it harder. Start a business."',"Humor","Sí, molde"),(1_600_000,'"Life\'s hard... try editing a video for 3 hours to get 10 likes" (collab @memequeen)',"Humor","Sí, molde")]),
("renee_warren",97800,[(17_000_000,'"This was the best investment I ever made…" (collab con Dan Martell)',"Personal","Parcial"),(4_300_000,'"The best investment you can make is the person you choose to marry" (podcast)',"Personal","No"),(1_900_000,'"Cuando mi marido \'se olvida\' de que la cámara está grabando"',"Humor","Parcial")]),
("iamnatalie",415000,[(3_200_000,'"Footage de mí intentando llevar un negocio, ser mamá presente, tomar suplementos..." (barriendo el mar)',"Humor","Sí, molde"),(1_100_000,'"Don\'t underestimate how quickly you can change your reality" (venta)',"Venta","No"),(510_000,'"We\'re leaving!!!" (mudanza familiar)',"Personal","No")]),
("itsamygolding",29000,[(317_000,'Lanzamiento de su proyecto para mamás ("Big girl pants on… join this revolution")',"Personal","Parcial"),(128_000,"Collab con @louisetruman: agradecimiento por su recorrido","Personal","No"),(41_600,"Collab @louisetruman: su experiencia en Dragons' Den","Personal","Parcial")]),
("cadacamino",13000,[(745_000,'"Cuando amas el proceso, trabajar por una meta se disfruta el doble"',"Opinión","Parcial"),(196_000,'"Mindset reminder": miedo a mostrarte y usar tu voz',"Opinión","Sí, molde"),(160_000,'"Salite del molde si te empieza a apretar"',"Opinión","Parcial")]),
("adrianabubori",114000,[(4_300_000,'Video hecho con IA del "nuevo Tesla Roadster" volando en NY',"Educativo","No"),(4_100_000,'"READ THIS — siempre soñé con mudarme al exterior y trabajar a mi manera"',"Personal","Sí, molde"),(2_900_000,'Misma historia "READ THIS" (versión anterior)',"Personal","Sí, molde")]),
("itsmeganlieu",241000,[(15_600_000,'"Reveal" de dejar las extensiones de pestañas ("lashless life")',"Personal","No"),(12_100_000,"Claude Code vs ChatGPT para programar","Pago","Parcial"),(10_100_000,"Reacción de su papá a un Porsche Macan EV nuevo","Personal","No")]),
("ai_herway",68400,[(7_000_000,'"\'Qué suerte que tu marido es exitoso\'": por qué está en business class (única mujer data scientist)',"Personal","Sí, molde"),(226_000,'Secuela del video de business class "con más pruebas"',"Personal","Sí, molde"),(91_700,"Su paso por ABC Radio: el riesgo silencioso de la IA para las mujeres","Opinión","Parcial")]),
("polinabatee",63400,[(2_700_000,'"Quizás estás publicando tus Reels mal"',"Educativo","Sí, molde"),(2_100_000,"Mismo tema, versión repost","Educativo","Sí, molde"),(451_000,'Hack de Instagram paso a paso con CTA "comment HACK"',"Educativo","Parcial")]),
("angcatalan",3821,[(123_000,'"Cada problema de negocio tiene un dueño en marketing"',"Educativo","Parcial"),(14_800,'"Las palabras que usás determinan si te compran"',"Educativo","Parcial"),(10_300,"Lista de libros para construir software","Educativo","No")]),
("byaditimishra",239000,[(4_200_000,'Lista gratis de sitios para trabajo remoto ("comment REMOTE")',"Educativo","Parcial"),(1_500_000,"Misma lista (versión anterior)","Educativo","Parcial"),(1_100_000,'Herramientas de IA para búsqueda laboral ("Not sponsored")',"Educativo","No")]),
("jeanelleteves",153000,[(14_400_000,'"As eldest daughter I make the rules": almuerzo semanal con su hermano',"Humor","Sí, molde"),(4_300_000,"#AwayPartner: vestidos que empacó para Cabo","Pago","No"),(1_600_000,'"Lunch with my brother Pt. 2"',"Humor","Sí, molde")]),
("nicolejizhar",21400,[(11_000_000,'"Eldest daughter mode activated" (humor hija mayor)',"Humor","Sí, molde"),(10_500_000,'Cumple 31: "no more \'next year this…\'", vivir el presente',"Personal","Sí, molde"),(6_100_000,'"The golden combo": hija mayor + hijo menor (pareja)',"Humor","Parcial")]),
("belindafurneauxharris",18600,[(47_500,"Lanzamiento de waitlist para retiro/workshop","Venta","No"),(34_800,'"Construir una marca no es solo para los conocidos"',"Opinión","Parcial"),(27_600,"Trabajo con diseñadoras de interiores","Educativo","No")]),
("itsvicchang",36200,[(349_000,'"Kidulting": LEGO pasó de 5% a ~45% de ventas a adultos',"Opinión","Sí, molde"),(170_000,'"Brands you should know: relojes"',"Opinión","Parcial"),(170_000,'Oportunidad en "muebles para gatos" (+45% búsquedas)',"Opinión","Parcial")]),
("mirandadoesbrands",54600,[(3_100_000,"¿Hacerse las uñas dejó de ser símbolo de estatus?","Opinión","Sí, molde"),(1_800_000,'Investigación de competencia con Similarweb ("comment LINK")',"Pago","Parcial"),(1_500_000,'De la era Mad Men al tropo de la "marketing girlie"',"Opinión","Sí, molde")]),
("nataliedawson",775000,[(7_000_000,'"Llamé al negocio de mi cliente en medio de una presentación"',"Personal","Sí, molde"),(5_900_000,'En The Diary of a CEO: "¿despedirías a tus empleados si hicieran trampa?"',"Opinión","Parcial"),(3_700_000,'"Si llamara a tu negocio ahora, ¿alguien atendería?"',"Opinión","Sí, molde")]),
("itshannahzhang",277000,[(13_700_000,"ChatGPT Work: convertir tu expertise en un servicio","Pago","Parcial"),(12_000_000,"Collab con @techwitharia para ChatGPT Work","Pago","No"),(11_000_000,'Partnership con LinkedIn: qué buscar cuando tu carrera es "no lineal"',"Pago","Sí, molde")]),
("marialeden",5241,[(202_000,'"LinkedIn no tiene la mejor reputación… pero un CV solo no alcanza"',"Opinión","Sí, molde"),(54_200,'Estético motivacional ("it does feel impossible sometimes but should be simple")',"Opinión","No"),(34_200,'"Step 1: pasá más tiempo con vos misma"',"Opinión","No")]),
("mollystovold",17600,[(9_100_000,"Historia de su embarazo: por qué separa negocio y vida","Personal","Parcial"),(2_800_000,'Relato de su parto ("my birth was a complete shambles")',"Personal","No"),(1_700_000,'Su longboard "unicornio" (surf)',"Personal","No")]),
("jamieegabrielle",189000,[(17_000_000,"Humor: grandes planes de contenido en vacaciones y un año sin publicar nada","Humor","Sí, molde"),(8_200_000,"Paseo en GMC HUMMER EV","Pago","No"),(6_700_000,"Organización de skincare con código de descuento (d'Alba)","Pago","No")]),
("melyssa_griffin",59200,[(67_300,"Episodio del BossBabe Podcast con ella","Personal","No"),(11_700,'"Un año con mi Violet" (su mascota)',"Personal","No"),(11_500,"Reflexión: querer controlar la vida es contraintuitivo","Opinión","No")]),
]

MOLDES = [
("Jeanelle Teves / Nicole Jizhar (14,4M y 11M)","\"As eldest daughter I make the rules\" / \"Eldest daughter mode activated\"","\"Modo ejecutiva activado\" · \"Como directora, yo pongo las reglas\" (misma plantilla de humor de identidad, con tu escena de oficina o familia)"),
("Jamie Gabrielle (17M)","Los grandes planes de contenido en vacaciones vs. un año sin publicar","\"Mis planes de publicar en LinkedIn en enero\" vs. \"yo en diciembre sin un solo post\""),
("Brit Crowther (2,2M y 1,6M)","\"Never in my likes, but always in my views\" · \"Life's hard. Make it harder. Start a business.\"","\"Nunca en mis likes, pero siempre en mis vistas\" (sobre tu jefe o tu CEO mirando tu LinkedIn) · \"La vida es difícil. Hacela más difícil: publicá en LinkedIn siendo directora.\""),
("Natalie Ellis (3,2M)","\"Footage de mí intentando llevar un negocio, ser mamá presente, tomar suplementos...\" (barriendo el mar)","\"Footage de mí intentando ser directora, tener voz en LinkedIn y contestar el grupo del colegio\" (misma imagen absurda)"),
("Nici Sweaney (7M)","\"'Qué suerte que tu marido es exitoso'\": por qué está en business class","\"'Qué suerte que te invitaron'\": por qué yo era la speaker (misma estructura: frase que te dijeron + la verdad)"),
("Sakshi Tyagi (2M)","\"How to Aura-Farm on LinkedIn\"","\"Cómo hacer aura farming en LinkedIn (versión ejecutiva)\": dejar el término en inglés, que es tendencia"),
("Polina Batishcheva (2,7M)","\"Quizás estás publicando tus Reels mal\"","\"Quizás estás publicando en LinkedIn mal\""),
("Diandra Escobar (354K, en español)","Aparezco en un ranking junto a presidentes de empresas gracias a LinkedIn","\"Salí #1 del mundo en Favikon, al lado de CEOs globales, gracias a LinkedIn\""),
("Nicole Jizhar (10,5M)","Cumple 31: \"no more 'next year this…'\"","\"No más 'el año que viene empiezo a publicar en LinkedIn'\""),
("Natalie Dawson (7M y 3,7M)","\"Llamé al negocio de mi cliente en medio de una presentación\" · \"Si llamara a tu negocio ahora, ¿alguien atendería?\"","\"Busqué en Google a cada persona del panel en medio del evento\" · \"Si un headhunter buscara tu nombre ahora, ¿qué encontraría?\""),
("Miranda Shanahan (3,1M)","\"¿Hacerse las uñas dejó de ser símbolo de estatus?\"","\"¿Tener cargo dejó de ser símbolo de estatus? (Ahora lo es tener voz)\""),
("Vic Chang (349K)","\"Kidulting\": dato sorprendente + tendencia con nombre","Un dato con nombre propio: \"El 'quiet posting': ejecutivas que leen todo y no publican nada\""),
("Adriana Bubori (4,1M)","\"READ THIS — siempre soñé con…\" (texto en pantalla sobre video de b-roll)","\"LEÉ ESTO: veinte años de carrera y me daba vergüenza escribir tres líneas en LinkedIn\""),
("Maria Leden (202K)","\"LinkedIn no tiene la mejor reputación… pero un CV solo no alcanza\"","\"LinkedIn no tiene la mejor reputación… pero un cargo solo no alcanza\""),
("Cada Camino (196K, en español)","\"Mindset reminder\": miedo a mostrarte y usar tu voz","\"Recordatorio para ejecutivas\": miedo a mostrarte y usar tu voz (casi igual; ya está en tu idioma y en tu tema)"),
]

def fmt(n):
    if n>=1_000_000: return f"{n/1_000_000:.1f}".replace(".0","").replace(".",",")+" M"
    if n>=1000: return f"{n/1000:.1f}".replace(".0","").replace(".",",")+" K"
    return str(n)

rows=[]
for cuenta,seg,reels in D:
    for v,desc,tipo,sirve in reels:
        rows.append(dict(cuenta=cuenta,seg=seg,v=v,desc=desc,tipo=tipo,sirve=sirve,ratio=v/seg))
rows.sort(key=lambda r:-r["ratio"])

doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=sec.page_height,sec.page_width
for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(sec,m,Cm(1.6))
st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10)

def shade(cell,hexcolor):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement("w:shd")
    s.set(qn("w:val"),"clear"); s.set(qn("w:color"),"auto"); s.set(qn("w:fill"),hexcolor); tcPr.append(s)

def setw(t,ws):
    t.autofit=False
    tblPr=t._tbl.tblPr; lay=OxmlElement("w:tblLayout"); lay.set(qn("w:type"),"fixed"); tblPr.append(lay)
    for j,w in enumerate(ws):
        t.columns[j].width=Cm(w)
        for c in t.columns[j].cells: c.width=Cm(w)

doc.add_heading("Los reels más virales de las 29 referentes",0)
p=doc.add_paragraph("Fuente: auditoría con Claude en Chrome, 22/10/2026 (los 3 reels más vistos de cada cuenta). "
 "Ordenados por cuánto viajó cada reel respecto del tamaño de su cuenta (vistas ÷ seguidores): "
 "cuanto más alto, más prueba hay de que el formato funciona solo, sin depender de la audiencia. Eso es lo que necesitamos al inicio.")
doc.add_paragraph("Límites: la auditoría no guardó los links ni el caption completo. Para tener el texto exacto de cada uno, "
 "bajar el Excel de Favikon de la cuenta. Los reels marcados como Pago pueden tener vistas infladas porque la marca pagó para amplificarlos; "
 "en cuentas grandes, el top sale de los reels más recientes revisados, no de todo el historial.")
doc.add_paragraph("Columna \"¿Sirve?\": Sí, molde = el gancho o la plantilla se puede copiar casi literal en tu nicho · "
 "Parcial = sirve la idea, no la plantilla · No = depende de su vida, de una marca o de otro nicho.")

doc.add_heading("1. Los moldes para copiar (cambiando lo mínimo)",1)
t=doc.add_table(rows=1,cols=3); t.style="Light Grid Accent 1"
for i,h in enumerate(["De quién","Gancho original","Tu versión"]): t.rows[0].cells[i].text=h
for a,b,c in MOLDES:
    r=t.add_row().cells; r[0].text=a; r[1].text=b; r[2].text=c
setw(t,[5,9.5,11])

doc.add_heading(f"2. Los {len(rows)} reels, de más a menos viaje",1)
t=doc.add_table(rows=1,cols=8); t.style="Light Grid Accent 1"
heads=["#","Cuenta","Seguidores","Vistas","Viajó (x seguidores)","De qué se trata / gancho","Tipo","¿Sirve?"]
for i,h in enumerate(heads): t.rows[0].cells[i].text=h
colors={"Sí, molde":"D9EAD3","Parcial":"FFF2CC","No":"F4F4F4"}
for i,r in enumerate(rows,1):
    c=t.add_row().cells
    vals=[str(i),"@"+r["cuenta"],fmt(r["seg"]),fmt(r["v"]),f"{r['ratio']:,.0f}x".replace(",","."),r["desc"],r["tipo"],r["sirve"]]
    for j,v in enumerate(vals): c[j].text=v
    shade(c[7],colors[r["sirve"]])
setw(t,[0.9,3.4,2,1.8,2.2,10.6,1.8,2])

doc.add_heading("3. Lo que muestran estos 87 reels",1)
n_si=sum(r['sirve']=="Sí, molde" for r in rows)
tipos={}
for r in rows: tipos[r["tipo"]]=tipos.get(r["tipo"],0)+1
top20=rows[:20]; t20={}
for r in top20: t20[r["tipo"]]=t20.get(r["tipo"],0)+1
for b in [
 f"De los {len(rows)} reels, {n_si} tienen un gancho o una plantilla que se puede copiar casi literal en tu nicho.",
 "Entre los 20 que más viajaron: " + ", ".join(f"{k} {v}" for k,v in sorted(t20.items(), key=lambda x:-x[1])) + ". Lo que viaja solo es humor de identidad e historia personal; lo educativo viaja menos.",
 "Las cuentas chicas con reels gigantes (Nicole Jizhar, Molly Stovold, Brit Crowther, Sakshi Tyagi, Nici Sweaney) son la prueba de que el formato, y no el tamaño, hace el alcance.",
 "Los formatos que más se repiten: identidad (\"modo X activado\", \"como hija mayor…\"), frase que te dijeron + la verdad, meme de frase en pantalla, \"quizás estás haciendo X mal\", y opinión sobre una tendencia con nombre propio.",
]:
    doc.add_paragraph(b,style="List Bullet")

doc.save("/home/user/Claude-VL/instagram/reels-virales-referentes.docx")
print(len(rows),"reels;",n_si,"molde; top20:",t20)
